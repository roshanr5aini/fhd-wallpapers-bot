import os
import re
import io
import asyncio
import threading
import requests

from flask import Flask, request

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
    Table,
    TableStyle,
    KeepTogether,
)
from reportlab.lib import colors
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from docx import Document as DocxDocument


BOT_TOKEN = os.getenv("BOT_TOKEN")
SERPAPI_KEY = os.getenv("SERPAPI_KEY")
REMOVE_BG_API_KEY = os.getenv("REMOVE_BG_API_KEY")

TIMEOUT = 20
RESULTS_PER_PAGE = 5

USER_SEARCHES = {}
USER_MODES = {}

# Text -> PDF sessions.
# Each user can paste multiple messages before pressing Generate PDF.
USER_PDF_SESSIONS = {}


# ==========================================
# QUERY PARSER
# ==========================================

def parse_query(text):
    original = text.strip()
    lower = original.lower()

    device = "any"
    quality = "normal"

    mobile_words = [
        "mobile", "phone", "smartphone", "iphone",
        "android", "portrait", "vertical",
    ]

    desktop_words = [
        "laptop", "desktop", "pc", "computer",
        "monitor", "landscape", "wide",
    ]

    quality_words = ["4k", "uhd", "2160p", "ultra hd"]

    if any(re.search(r"\b" + re.escape(word) + r"\b", lower)
           for word in mobile_words):
        device = "mobile"
    elif any(re.search(r"\b" + re.escape(word) + r"\b", lower)
             for word in desktop_words):
        device = "desktop"

    if any(word in lower for word in quality_words):
        quality = "4k"

    remove_words = (
        mobile_words + desktop_words + quality_words +
        [
            "wallpaper", "wallpapers", "wall", "background",
            "backgrounds", "hd", "fhd", "full hd", "for",
            "me", "please",
        ]
    )

    query = original
    for word in sorted(remove_words, key=len, reverse=True):
        query = re.sub(
            r"\b" + re.escape(word) + r"\b",
            " ",
            query,
            flags=re.IGNORECASE,
        )

    query = re.sub(r"\s+", " ", query).strip()
    return query, device, quality


# ==========================================
# IMAGE QUALITY / RATIO FILTER
# ==========================================

def suitable_image(item, device, quality):
    width = item.get("original_width")
    height = item.get("original_height")

    try:
        width = int(width)
        height = int(height)
    except Exception:
        return False

    if width <= 0 or height <= 0:
        return False

    ratio = width / height

    if device == "mobile":
        if height <= width:
            return False
        if ratio < 0.40 or ratio > 0.80:
            return False
        if height < 1200:
            return False
        if quality == "4k" and height < 1800:
            return False
        return True

    if device == "desktop":
        if width <= height:
            return False
        if ratio < 1.30 or ratio > 2.40:
            return False
        if width < 1400:
            return False
        if quality == "4k" and width < 2500:
            return False
        return True

    if min(width, height) < 900:
        return False

    if quality == "4k" and max(width, height) < 2500:
        return False

    return True


# ==========================================
# PNG IMAGE FILTER
# ==========================================

def suitable_png(item):
    width = item.get("original_width")
    height = item.get("original_height")

    try:
        width = int(width)
        height = int(height)
    except Exception:
        return False

    if width <= 0 or height <= 0:
        return False

    if max(width, height) < 500:
        return False

    combined = " ".join([
        str(item.get("original", "")),
        str(item.get("title", "")),
        str(item.get("source", "")),
        str(item.get("link", "")),
    ]).lower()

    png_signals = [
        ".png", "png", "transparent",
        "no background", "cutout", "render",
    ]

    if any(word in combined for word in png_signals):
        return True

    return min(width, height) >= 700


# ==========================================
# GOOGLE IMAGES SEARCH
# ==========================================

def google_images_search(query, device, quality, page):
    search_query = query + " wallpaper"

    if device == "mobile":
        search_query += " mobile portrait"
    elif device == "desktop":
        search_query += " desktop landscape"

    if quality == "4k":
        search_query += " 4K"

    params = {
        "engine": "google_images",
        "q": search_query,
        "api_key": SERPAPI_KEY,
        "safe": "active",
        "hl": "en",
        "gl": "us",
        "ijn": page,
    }

    try:
        response = requests.get(
            "https://serpapi.com/search",
            params=params,
            timeout=TIMEOUT,
        )
        if response.status_code != 200:
            print("SerpApi HTTP:", response.status_code)
            return []
        return response.json().get("images_results", [])
    except Exception as e:
        print("SerpApi error:", e)
        return []


def google_png_search(query, page):
    search_query = query + " PNG transparent background"

    params = {
        "engine": "google_images",
        "q": search_query,
        "api_key": SERPAPI_KEY,
        "safe": "active",
        "hl": "en",
        "gl": "us",
        "ijn": page,
    }

    try:
        response = requests.get(
            "https://serpapi.com/search",
            params=params,
            timeout=TIMEOUT,
        )
        if response.status_code != 200:
            print("SerpApi PNG HTTP:", response.status_code)
            return []
        return response.json().get("images_results", [])
    except Exception as e:
        print("SerpApi PNG error:", e)
        return []


def collect_results(query, device, quality, start_page=0):
    collected = []
    for page in range(start_page, start_page + 3):
        for item in google_images_search(query, device, quality, page):
            if item.get("original") and suitable_image(item, device, quality):
                collected.append(item)
    return collected


def collect_png_results(query, start_page=0):
    collected = []
    for page in range(start_page, start_page + 3):
        for item in google_png_search(query, page):
            if item.get("original") and suitable_png(item):
                collected.append(item)
    return collected


def remove_duplicates(items, seen):
    final = []
    local_seen = set(seen)

    for item in items:
        url = item.get("original")
        if not url:
            continue

        key = url.split("?")[0].lower().strip()
        if key in local_seen:
            continue

        local_seen.add(key)
        final.append(item)

    return final


# ==========================================
# REMOVE BACKGROUND
# ==========================================

def remove_background(image_bytes):
    if not REMOVE_BG_API_KEY:
        return None, "API key is missing."

    try:
        response = requests.post(
            "https://api.remove.bg/v1.0/removebg",
            files={
                "image_file": (
                    "image.jpg",
                    image_bytes,
                    "image/jpeg",
                )
            },
            data={"size": "auto", "format": "png"},
            headers={"X-Api-Key": REMOVE_BG_API_KEY},
            timeout=60,
        )

        if response.status_code == 200:
            return response.content, None

        print(
            "Remove.bg HTTP:",
            response.status_code,
            response.text[:500],
        )
        return None, f"Remove.bg error {response.status_code}"

    except Exception as e:
        print("Remove.bg request error:", e)
        return None, str(e)


# ==========================================
# TEXT -> PDF
# ==========================================

def clean_pdf_text(text):
    # Normalize common line endings and invisible characters.
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", "")
    return text.strip()


def escape_pdf_text(text):
    # ReportLab Paragraph uses a small XML-like markup language.
    from xml.sax.saxutils import escape
    return escape(text, {"'": "&apos;"})


def add_page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.drawCentredString(
        A4[0] / 2,
        10 * mm,
        f"Page {doc.page}",
    )
    canvas.restoreState()


def make_pdf_from_text(text):
    text = clean_pdf_text(text)
    if not text:
        raise ValueError("There is no text to convert.")

    output = io.BytesIO()

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "PDFBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=16,
        spaceAfter=8,
        alignment=TA_LEFT,
    )
    heading = ParagraphStyle(
        "PDFHeading",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=21,
        spaceBefore=8,
        spaceAfter=10,
    )
    subheading = ParagraphStyle(
        "PDFSubHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        spaceBefore=7,
        spaceAfter=7,
    )
    bullet = ParagraphStyle(
        "PDFBullet",
        parent=body,
        leftIndent=14,
        firstLineIndent=-8,
        spaceAfter=5,
    )

    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="FHD Wallpapers Bot - PDF",
    )

    story = []

    lines = text.split("\n")
    paragraph_lines = []

    def flush_paragraph():
        nonlocal paragraph_lines
        if not paragraph_lines:
            return

        paragraph = " ".join(
            line.strip() for line in paragraph_lines
        ).strip()

        if paragraph:
            story.append(
                Paragraph(
                    escape_pdf_text(paragraph),
                    body,
                )
            )

        paragraph_lines = []

    for raw in lines:
        line = raw.strip()

        if not line:
            flush_paragraph()
            story.append(Spacer(1, 2))
            continue

        # Explicit page-break marker.
        if line.lower() in {
            "[page break]",
            "--- page break ---",
            "=== page break ===",
        }:
            flush_paragraph()
            story.append(PageBreak())
            continue

        # Markdown-like headings.
        if line.startswith("### "):
            flush_paragraph()
            story.append(
                Paragraph(
                    escape_pdf_text(line[4:].strip()),
                    subheading,
                )
            )
            continue

        if line.startswith("## "):
            flush_paragraph()
            story.append(
                Paragraph(
                    escape_pdf_text(line[3:].strip()),
                    heading,
                )
            )
            continue

        if line.startswith("# "):
            flush_paragraph()
            story.append(
                Paragraph(
                    escape_pdf_text(line[2:].strip()),
                    heading,
                )
            )
            continue

        # Common list formats.
        if re.match(r"^[-*•]\s+", line):
            flush_paragraph()
            item_text = re.sub(r"^[-*•]\s+", "", line)
            story.append(
                Paragraph(
                    "• " + escape_pdf_text(item_text),
                    bullet,
                )
            )
            continue

        if re.match(r"^\d+[.)]\s+", line):
            flush_paragraph()
            story.append(
                Paragraph(
                    escape_pdf_text(line),
                    bullet,
                )
            )
            continue

        # Treat short all-caps lines as headings.
        if (
            len(line) <= 90
            and len(line.split()) <= 12
            and line.upper() == line
            and re.search(r"[A-Z]", line)
        ):
            flush_paragraph()
            story.append(
                Paragraph(
                    escape_pdf_text(line),
                    subheading,
                )
            )
            continue

        paragraph_lines.append(line)

    flush_paragraph()

    if not story:
        story.append(Paragraph("Empty document", body))

    doc.build(
        story,
        onFirstPage=add_page_number,
        onLaterPages=add_page_number,
    )

    output.seek(0)
    return output


def extract_docx_text(file_bytes):
    document = DocxDocument(io.BytesIO(file_bytes))
    parts = []

    # Paragraphs preserve basic Word document reading order.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            parts.append("")
            continue

        style_name = (paragraph.style.name or "").lower()

        if "title" in style_name:
            parts.append("# " + text)
        elif "heading 1" in style_name:
            parts.append("## " + text)
        elif "heading 2" in style_name or "heading 3" in style_name:
            parts.append("### " + text)
        elif "list" in style_name:
            parts.append("• " + text)
        else:
            parts.append(text)

    # Add tables in a simple readable text format.
    for table in document.tables:
        parts.append("")
        for row in table.rows:
            cells = [
                cell.text.replace("\n", " ").strip()
                for cell in row.cells
            ]
            parts.append(" | ".join(cells))
        parts.append("")

    return "\n".join(parts).strip()


def pdf_menu_keyboard():
    return InlineKeyboardMarkup([
        [
            InlineKeyboardButton(
                "✅ Generate PDF",
                callback_data="pdf_generate",
            ),
            InlineKeyboardButton(
                "🗑️ Clear",
                callback_data="pdf_clear",
            ),
        ],
        [
            InlineKeyboardButton(
                "🏠 Main Menu",
                callback_data="main_menu",
            )
        ],
    ])


def pdf_mode_message():
    return (
        "📄 Text → PDF Mode\n\n"
        "Send your content in any of these ways:\n\n"
        "📝 Paste text directly here\n"
        "📄 Upload a .txt file\n"
        "📘 Upload a .docx Word file\n\n"
        "You can send multiple text messages too.\n"
        "When finished, tap ✅ Generate PDF."
    )


# ==========================================
# SEND NEXT 5 WALLPAPERS
# ==========================================

async def send_five(context, user_id, chat_id):
    data = USER_SEARCHES.get(user_id)
    if not data:
        return

    query = data["query"]
    device = data["device"]
    quality = data["quality"]
    results = data["results"]
    index = data["index"]
    sent = 0

    while sent < RESULTS_PER_PAGE:
        if index >= len(results):
            next_page = data["next_page"]
            new_results = collect_results(
                query, device, quality, next_page
            )
            new_results = remove_duplicates(
                new_results, data["seen"]
            )
            data["next_page"] += 3
            results.extend(new_results)

            if not new_results:
                break

        if index >= len(results):
            break

        item = results[index]
        index += 1
        url = item.get("original")

        if not url:
            continue

        key = url.split("?")[0].lower().strip()
        if key in data["seen"]:
            continue

        try:
            width = item.get("original_width", "?")
            height = item.get("original_height", "?")
            caption = f"🖼️ {query}\n📐 {width}×{height}"

            await context.bot.send_photo(
                chat_id=chat_id,
                photo=url,
                caption=caption,
            )

            data["seen"].add(key)
            sent += 1

        except Exception as e:
            print("Telegram image error:", e)
            continue

    data["index"] = index

    if index < len(results) or data["next_page"] < 12:
        keyboard = [
            [
                InlineKeyboardButton(
                    "➕ More 5",
                    callback_data="more5",
                )
            ],
            [
                InlineKeyboardButton(
                    "🏠 Main Menu",
                    callback_data="main_menu",
                )
            ],
        ]
        await context.bot.send_message(
            chat_id=chat_id,
            text="Want to explore more wallpapers?",
            reply_markup=InlineKeyboardMarkup(keyboard),
        )
    else:
        await context.bot.send_message(
            chat_id=chat_id,
            text="ℹ️ No more suitable results are available.",
        )


# ==========================================
# SEND NEXT 5 PNGs
# ==========================================

async def send_png_five(context, user_id, chat_id):
    data = USER_SEARCHES.get(user_id)
    if not data:
        return

    query = data["query"]
    results = data["results"]
    index = data["index"]
    sent = 0

    while sent < RESULTS_PER_PAGE:
        if index >= len(results):
            next_page = data["next_page"]
            new_results = collect_png_results(query, next_page)
            new_results = remove_duplicates(
                new_results, data["seen"]
            )
            data["next_page"] += 3
            results.extend(new_results)

            if not new_results:
                break

        if index >= len(results):
            break

        item = results[index]
        index += 1
        url = item.get("original")

        if not url:
            continue

        key = url.split("?")[0].lower().strip()
        if key in data["seen"]:
            continue

        try:
            width = item.get("original_width", "?")
            height = item.get("original_height", "?")

            await context.bot.send_document(
                chat_id=chat_id,
                document=url,
                caption=(
                    f"🖼️ PNG: {query}\n"
                    f"📐 {width}×{height}\n"
                    f"✨ Transparent PNG search"
                ),
            )

            data["seen"].add(key)
            sent += 1

        except Exception as e:
            print("Telegram PNG error:", e)

            try:
                await context.bot.send_photo(
                    chat_id=chat_id,
                    photo=url,
                    caption=(
                        f"🖼️ PNG: {query}\n"
                        f"📐 {width}×{height}"
                    ),
                )
                data["seen"].add(key)
                sent += 1
            except Exception as e2:
                print("Telegram PNG photo error:", e2)
                continue

    data["index"] = index

    if index < len(results) or data["next_page"] < 12:
        keyboard = [
            [
                InlineKeyboardButton(
                    "➕ More 5",
                    callback_data="png_more5",
                )
            ],
            [
                InlineKeyboardButton(
                    "🏠 Main Menu",
                    callback_data="main_menu",
                )
            ],
        ]
        await context.bot.send_message(
            chat_id=chat_id,
            text="Want more PNG images?",
            reply_markup=InlineKeyboardMarkup(keyboard),
        )
    else:
        await context.bot.send_message(
            chat_id=chat_id,
            text="ℹ️ No more suitable PNG results are available.",
        )


# ==========================================
# MAIN MENU
# ==========================================

async def show_main_menu(update, context, edit=False):
    keyboard = [
        [
            InlineKeyboardButton(
                "🖼️ Wallpaper Search",
                callback_data="wallpaper_mode",
            )
        ],
        [
            InlineKeyboardButton(
                "🖼️ PNG Images",
                callback_data="png_mode",
            )
        ],
        [
            InlineKeyboardButton(
                "✂️ Remove Background",
                callback_data="remove_bg_mode",
            )
        ],
        [
            InlineKeyboardButton(
                "📄 Text → PDF",
                callback_data="pdf_mode",
            )
        ],
    ]

    text = (
        "✨ FHD Wallpapers Bot\n\n"
        "Choose what you want to use:\n\n"
        "🖼️ Wallpaper Search\n"
        "Search HD & 4K wallpapers.\n\n"
        "🖼️ PNG Images\n"
        "Search PNG & transparent images.\n\n"
        "✂️ Remove Background\n"
        "Remove photo background and get a transparent PNG.\n\n"
        "📄 Text → PDF\n"
        "Convert pasted text, .txt or .docx into a clean PDF."
    )

    if edit:
        await update.callback_query.message.edit_text(
            text,
            reply_markup=InlineKeyboardMarkup(keyboard),
        )
    else:
        await update.message.reply_text(
            text,
            reply_markup=InlineKeyboardMarkup(keyboard),
        )


# ==========================================
# START
# ==========================================

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id

    USER_MODES[user_id] = "wallpaper"
    USER_SEARCHES.pop(user_id, None)
    USER_PDF_SESSIONS.pop(user_id, None)

    await show_main_menu(update, context)


# ==========================================
# REMOVE BACKGROUND PHOTO HANDLER
# ==========================================

async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    chat_id = update.effective_chat.id

    mode = USER_MODES.get(user_id, "wallpaper")

    if mode != "remove_bg":
        return

    status = await update.message.reply_text(
        "✂️ Removing background...\n\n"
        "⏳ Please wait."
    )

    try:
        photo = update.message.photo[-1]
        telegram_file = await context.bot.get_file(photo.file_id)
        image_bytes = await telegram_file.download_as_bytearray()

        result, error = remove_background(bytes(image_bytes))

        if not result:
            await status.edit_text(
                "❌ Background removal failed.\n\n"
                f"Reason: {error}"
            )
            return

        await status.edit_text(
            "✅ Background removed!\n\n"
            "📤 Sending transparent PNG..."
        )

        await context.bot.send_document(
            chat_id=chat_id,
            document=io.BytesIO(result),
            filename="no_background.png",
            caption="✂️ Background Removed\n✨ Transparent PNG",
        )

        await status.delete()

        keyboard = [
            [
                InlineKeyboardButton(
                    "✂️ Remove Another",
                    callback_data="remove_bg_mode",
                )
            ],
            [
                InlineKeyboardButton(
                    "🏠 Main Menu",
                    callback_data="main_menu",
                )
            ],
        ]

        await context.bot.send_message(
            chat_id=chat_id,
            text="What would you like to do next?",
            reply_markup=InlineKeyboardMarkup(keyboard),
        )

    except Exception as e:
        print("Remove background handler error:", e)
        await status.edit_text(
            "❌ Something went wrong.\n\n"
            "Please try another photo."
        )


# ==========================================
# TEXT/DOCUMENT -> PDF HANDLER
# ==========================================

async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    chat_id = update.effective_chat.id

    if USER_MODES.get(user_id) != "pdf":
        return

    document = update.message.document
    filename = (document.file_name or "").lower()

    if not (filename.endswith(".txt") or filename.endswith(".docx")):
        await update.message.reply_text(
            "❌ Please upload only a .txt or .docx file."
        )
        return

    status = await update.message.reply_text(
        "📄 Reading your document...\n\n"
        "⏳ Please wait."
    )

    try:
        telegram_file = await context.bot.get_file(document.file_id)
        file_bytes = bytes(
            await telegram_file.download_as_bytearray()
        )

        if filename.endswith(".txt"):
            text = file_bytes.decode("utf-8-sig", errors="replace")
        else:
            text = extract_docx_text(file_bytes)

        text = clean_pdf_text(text)

        if not text:
            await status.edit_text(
                "❌ The uploaded file appears to be empty."
            )
            return

        USER_PDF_SESSIONS.setdefault(user_id, [])
        USER_PDF_SESSIONS[user_id].append(text)

        total_chars = sum(
            len(part) for part in USER_PDF_SESSIONS[user_id]
        )

        await status.edit_text(
            "✅ Document added to your PDF.\n\n"
            f"📄 File: {document.file_name}\n"
            f"🔤 Characters: {total_chars}\n\n"
            "You can upload/paste more content, or tap "
            "✅ Generate PDF.",
            reply_markup=pdf_menu_keyboard(),
        )

    except Exception as e:
        print("Document -> PDF error:", e)
        await status.edit_text(
            "❌ Could not read this file.\n\n"
            "Please try another .txt or .docx file."
        )


# ==========================================
# NORMAL TEXT MESSAGE
# ==========================================

async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()

    if not text:
        return

    user_id = update.effective_user.id
    chat_id = update.effective_chat.id
    mode = USER_MODES.get(user_id, "wallpaper")

    # ======================================
    # PDF MODE
    # ======================================

    if mode == "pdf":
        USER_PDF_SESSIONS.setdefault(user_id, [])
        USER_PDF_SESSIONS[user_id].append(text)

        total_chars = sum(
            len(part) for part in USER_PDF_SESSIONS[user_id]
        )

        await update.message.reply_text(
            "📝 Text added to your PDF.\n\n"
            f"🔤 Total characters: {total_chars}\n\n"
            "Send more text if needed, then tap "
            "✅ Generate PDF.",
            reply_markup=pdf_menu_keyboard(),
        )
        return

    # ======================================
    # PNG MODE
    # ======================================

    if mode == "png":
        query = re.sub(
            r"\bpng\b",
            " ",
            text,
            flags=re.IGNORECASE,
        )
        query = re.sub(r"\s+", " ", query).strip()

        if not query:
            await update.message.reply_text(
                "❌ Please enter what PNG you want.\n\n"
                "Example:\n"
                "• hand png\n"
                "• anime hair png\n"
                "• car png"
            )
            return

        status = await update.message.reply_text(
            f"🔎 Searching PNGs for: {query}\n\n"
            "⏳ Finding transparent images..."
        )

        USER_SEARCHES[user_id] = {
            "query": query,
            "device": "any",
            "quality": "normal",
            "results": [],
            "index": 0,
            "next_page": 0,
            "seen": set(),
            "type": "png",
        }

        data = USER_SEARCHES[user_id]
        results = remove_duplicates(
            collect_png_results(query, 0),
            set(),
        )
        data["results"] = results
        data["next_page"] = 3

        if not results:
            await status.edit_text(
                f"❌ No suitable PNGs found for '{query}'.\n\n"
                "Try another search."
            )
            USER_SEARCHES.pop(user_id, None)
            return

        await status.edit_text(
            f"✅ PNGs found for: {query}\n\n"
            "🖼️ Sending the best 5..."
        )

        await send_png_five(context, user_id, chat_id)
        return

    # ======================================
    # WALLPAPER MODE
    # ======================================

    query, device, quality = parse_query(text)

    if not query:
        await update.message.reply_text(
            "❌ Please enter a subject to search for."
        )
        return

    device_text = {
        "mobile": "📱 Mobile",
        "desktop": "💻 Desktop",
        "any": "🖼️ Any format",
    }[device]

    quality_text = {
        "normal": "✨ High Quality",
        "4k": "🔥 4K Preferred",
    }[quality]

    status = await update.message.reply_text(
        f"🔎 Searching for: {query}\n"
        f"{device_text}\n"
        f"{quality_text}\n\n"
        "⏳ Finding the best wallpapers..."
    )

    USER_SEARCHES[user_id] = {
        "query": query,
        "device": device,
        "quality": quality,
        "results": [],
        "index": 0,
        "next_page": 0,
        "seen": set(),
        "type": "wallpaper",
    }

    data = USER_SEARCHES[user_id]
    results = remove_duplicates(
        collect_results(query, device, quality, 0),
        set(),
    )

    data["results"] = results
    data["next_page"] = 3

    if not results:
        await status.edit_text(
            f"❌ No suitable wallpapers were found for '{query}'.\n\n"
            "Try another search term."
        )
        USER_SEARCHES.pop(user_id, None)
        return

    await status.edit_text(
        f"✅ Wallpapers found for: {query}\n\n"
        f"{device_text}\n"
        f"{quality_text}\n\n"
        "🖼️ Sending the best 5..."
    )

    await send_five(context, user_id, chat_id)


# ==========================================
# CALLBACK BUTTONS
# ==========================================

async def button_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()

    user_id = query.from_user.id
    chat_id = query.message.chat_id

    if query.data == "wallpaper_mode":
        USER_MODES[user_id] = "wallpaper"
        USER_PDF_SESSIONS.pop(user_id, None)

        await query.message.edit_text(
            "🖼️ Wallpaper Search Mode\n\n"
            "Now type anything you want to search for.\n\n"
            "Examples:\n"
            "• Ferrari 4K\n"
            "• Tokyo night mobile\n"
            "• Space wallpaper\n\n"
            "🚀 Send your search:"
        )
        return

    if query.data == "png_mode":
        USER_MODES[user_id] = "png"
        USER_PDF_SESSIONS.pop(user_id, None)

        await query.message.edit_text(
            "🖼️ PNG Search Mode\n\n"
            "Type anything you want as a PNG.\n\n"
            "Examples:\n"
            "• hand png\n"
            "• anime hair png\n"
            "• car png\n"
            "• Naruto png\n"
            "• flower png\n\n"
            "🚀 Send your PNG search:"
        )
        return

    if query.data == "remove_bg_mode":
        USER_MODES[user_id] = "remove_bg"
        USER_PDF_SESSIONS.pop(user_id, None)

        await query.message.edit_text(
            "✂️ Remove Background Mode\n\n"
            "Send me a photo and I will remove its background.\n\n"
            "✨ You will receive a transparent PNG.\n\n"
            "📸 Send your photo:"
        )
        return

    # --------------------------------------
    # Text -> PDF mode
    # --------------------------------------

    if query.data == "pdf_mode":
        USER_MODES[user_id] = "pdf"
        USER_SEARCHES.pop(user_id, None)
        USER_PDF_SESSIONS[user_id] = []

        await query.message.edit_text(
            pdf_mode_message(),
            reply_markup=pdf_menu_keyboard(),
        )
        return

    # --------------------------------------
    # Generate PDF
    # --------------------------------------

    if query.data == "pdf_generate":
        parts = USER_PDF_SESSIONS.get(user_id, [])

        if not parts:
            await query.message.reply_text(
                "❌ Your PDF is empty.\n\n"
                "Paste some text or upload a .txt/.docx file first."
            )
            return

        await query.message.edit_text(
            "📄 Generating your PDF...\n\n"
            "⏳ Please wait."
        )

        try:
            combined_text = "\n\n".join(parts)
            pdf_file = make_pdf_from_text(combined_text)

            await context.bot.send_document(
                chat_id=chat_id,
                document=pdf_file,
                filename="converted_text.pdf",
                caption=(
                    "✅ PDF Ready!\n\n"
                    "📄 Your text has been converted into a "
                    "cleanly formatted PDF."
                ),
            )

            USER_PDF_SESSIONS.pop(user_id, None)

            await context.bot.send_message(
                chat_id=chat_id,
                text="What would you like to do next?",
                reply_markup=InlineKeyboardMarkup([
                    [
                        InlineKeyboardButton(
                            "📄 Create Another PDF",
                            callback_data="pdf_mode",
                        )
                    ],
                    [
                        InlineKeyboardButton(
                            "🏠 Main Menu",
                            callback_data="main_menu",
                        )
                    ],
                ]),
            )

        except Exception as e:
            print("PDF generation error:", e)
            await query.message.edit_text(
                "❌ PDF generation failed.\n\n"
                "Please try again with simpler text."
            )
        return

    # --------------------------------------
    # Clear PDF
    # --------------------------------------

    if query.data == "pdf_clear":
        USER_PDF_SESSIONS[user_id] = []

        await query.message.edit_text(
            "🗑️ PDF content cleared.\n\n" +
            pdf_mode_message(),
            reply_markup=pdf_menu_keyboard(),
        )
        return

    # --------------------------------------
    # Main menu
    # --------------------------------------

    if query.data == "main_menu":
        USER_MODES[user_id] = "wallpaper"
        USER_SEARCHES.pop(user_id, None)
        USER_PDF_SESSIONS.pop(user_id, None)

        await show_main_menu(update, context, edit=True)
        return

    # --------------------------------------
    # More wallpapers
    # --------------------------------------

    if query.data == "more5":
        if user_id not in USER_SEARCHES:
            await query.message.reply_text(
                "⚠️ This search session has expired.\n\n"
                "Please start a new search."
            )
            return

        await query.answer("Finding 5 different wallpapers...")
        await send_five(context, user_id, chat_id)
        return

    # --------------------------------------
    # More PNGs
    # --------------------------------------

    if query.data == "png_more5":
        if user_id not in USER_SEARCHES:
            await query.message.reply_text(
                "⚠️ This search session has expired.\n\n"
                "Please start a new PNG search."
            )
            return

        await query.answer("Finding 5 different PNGs...")
        await send_png_five(context, user_id, chat_id)
        return


# ==========================================
# ERROR HANDLER
# ==========================================

async def error_handler(update, context):
    print("BOT ERROR:", context.error)


# ==========================================
# BOT / FLASK
# ==========================================

flask_app = Flask(__name__)
BOT_APP = None
BOT_LOOP = None


async def run_bot():
    global BOT_APP, BOT_LOOP

    BOT_LOOP = asyncio.get_running_loop()

    BOT_APP = (
        Application.builder()
        .token(BOT_TOKEN)
        .build()
    )

    BOT_APP.add_handler(
        CommandHandler("start", start)
    )

    BOT_APP.add_handler(
        MessageHandler(filters.PHOTO, handle_photo)
    )

    # Must be before generic text handling.
    BOT_APP.add_handler(
        MessageHandler(
            filters.Document.ALL,
            handle_document,
        )
    )

    BOT_APP.add_handler(
        MessageHandler(
            filters.TEXT & ~filters.COMMAND,
            handle_message,
        )
    )

    BOT_APP.add_handler(
        CallbackQueryHandler(button_handler)
    )

    BOT_APP.add_error_handler(error_handler)

    await BOT_APP.initialize()
    await BOT_APP.start()

    print("🔥 FHD Wallpapers Bot webhook mode started!")

    while True:
        await asyncio.sleep(3600)


def run_flask():
    port = int(os.getenv("PORT", "10000"))

    print(f"🌐 Web server running on port {port}")

    flask_app.run(
        host="0.0.0.0",
        port=port,
        debug=False,
        use_reloader=False,
    )


def main():
    if not BOT_TOKEN:
        print("❌ BOT_TOKEN is missing!")
        return

    if not SERPAPI_KEY:
        print("❌ SERPAPI_KEY is missing!")
        return

    threading.Thread(
        target=run_flask,
        daemon=True,
    ).start()

    asyncio.run(run_bot())


@flask_app.route("/")
def health():
    return "FHD Wallpapers Bot is running!", 200


@flask_app.route("/webhook", methods=["POST"])
def webhook():
    if BOT_APP is None:
        return "Bot not ready", 503

    try:
        data = request.get_json(force=True)
        update = Update.de_json(data, BOT_APP.bot)
    except Exception as e:
        print("Webhook parse error:", e)
        return "Bad update", 400

    try:
        asyncio.run_coroutine_threadsafe(
            BOT_APP.process_update(update),
            BOT_LOOP,
        )
    except Exception as e:
        print("Webhook processing error:", e)
        return "Processing error", 500

    return "OK", 200


if __name__ == "__main__":
    main()
